"""
矮小銀河5天体の全13ビン解析スライドをPPTに追加。

各天体の既存スライド（点源汚染チェックスライド）の直後に挿入：
  1. 全13ビン S/N スペクトル（bins13_<name>.png）
  2. 低エネルギースカイマップ（skymaps13_low_<name>.png）
  3. 高エネルギースカイマップ（skymaps13_high_<name>.png）
  4. テンプレートマップ（templates_fig14_<name>.png）
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

PPT_PATH = Path("PPT/slides_v3_detailed_fixed2_final.pptx  -  Repaired.pptx")
BASE     = Path(".")
prs      = Presentation(str(PPT_PATH))

W = prs.slide_width
H = prs.slide_height
C_BG    = RGBColor(0x05, 0x05, 0x1A)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE = RGBColor(0xFF, 0xCC, 0x00)
C_CYAN  = RGBColor(0x00, 0xFF, 0xFF)

# 各天体の設定: name, label, 挿入先（直後）のスライドタイトルの一部
DWARFS = [
    ("draco",      "Draco",       "Draco dSph — 点源汚染"),
    ("sculptor",   "Sculptor",    "Sculptor dSph — 点源汚染"),
    ("ursa_minor", "Ursa Minor",  "Ursa Minor dSph — 点源汚染"),
    ("segue1",     "Segue 1",     "Segue 1 — 点源汚染"),
    ("coma_ber",   "Coma Ber.",   "Coma Berenices — 点源汚染"),
]

# 各天体のBin6 S/N（表示用）
BIN6_SN = {
    "draco":      "-1.33",
    "sculptor":   "+1.27",
    "ursa_minor": "+1.61",
    "segue1":     "-0.38",
    "coma_ber":   "+4.04 (点源汚染)",
}


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def set_bg(sl):
    bg = sl.background; fill = bg.fill
    fill.solid(); fill.fore_color.rgb = C_BG

def add_title(sl, text, color=C_TITLE, size=22):
    tb = sl.shapes.add_textbox(Inches(0.15), Inches(0.05), W-Inches(0.3), Inches(0.52))
    tf = tb.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; run = p.add_run()
    run.text = text; run.font.size = Pt(size)
    run.font.color.rgb = color; run.font.bold = True

def add_sep(sl):
    box = sl.shapes.add_shape(1, Inches(0.15), Inches(0.60), W-Inches(0.3), Inches(0.018))
    box.fill.solid(); box.fill.fore_color.rgb = C_CYAN
    box.line.fill.background()

def add_note(sl, text, color=None):
    c = color or RGBColor(0xAA, 0xAA, 0xAA)
    tb = sl.shapes.add_textbox(Inches(0.15), H-Inches(0.38), W-Inches(0.3), Inches(0.35))
    tf = tb.text_frame; p = tf.paragraphs[0]; run = p.add_run()
    run.text = text; run.font.size = Pt(10); run.font.color.rgb = c

def make_13bin_slide(name, label):
    """全13ビン S/N スペクトルスライド"""
    sl = blank_slide(prs); set_bg(sl)
    sn6 = BIN6_SN.get(name, "?")
    add_title(sl, f"{label} — 全13ビン S/N スペクトル（Totani Fig.8/9 相当）")
    add_sep(sl)
    img = Path(f"data/figure-dwarfs/{name}/bins13_{name}.png")
    if img.exists():
        sl.shapes.add_picture(str(img), Inches(0.15), Inches(0.65),
                              width=W-Inches(0.3))
    add_note(sl, f"ON=2°, OFF=2°-5° | Bin6 (20.76 GeV) S/N = {sn6}σ | "
                 f"Bin10以上（>169 GeV）は光子数不足のため統計的信頼性低")
    return sl

def make_skymaps_low_slide(name, label):
    """低エネルギービン スカイマップスライド"""
    sl = blank_slide(prs); set_bg(sl)
    add_title(sl, f"{label} — スカイマップ 低エネルギー Bin1-6（1.5–28 GeV）")
    add_sep(sl)
    img = Path(f"data/figure-dwarfs/{name}/skymaps13_low_{name}.png")
    if img.exists():
        sl.shapes.add_picture(str(img), Inches(0.15), Inches(0.65),
                              width=W-Inches(0.3))
    add_note(sl, "白実線=ON領域(2°), 白破線=OFF領域(5°), 十字=天体中心 | 0.5°/pix")
    return sl

def make_skymaps_high_slide(name, label):
    """高エネルギービン スカイマップスライド"""
    sl = blank_slide(prs); set_bg(sl)
    add_title(sl, f"{label} — スカイマップ 高エネルギー Bin7-13（35–814 GeV）")
    add_sep(sl)
    img = Path(f"data/figure-dwarfs/{name}/skymaps13_high_{name}.png")
    if img.exists():
        sl.shapes.add_picture(str(img), Inches(0.15), Inches(0.65),
                              width=W-Inches(0.3))
    add_note(sl, "Bin10以上は光子数 ≤3 のため統計的有意性なし。高エネルギーでの非検出確認用。")
    return sl

def make_template_slide(name, label):
    """テンプレートマップスライド（Totani Fig.14 相当）"""
    sl = blank_slide(prs); set_bg(sl)
    add_title(sl, f"{label} — テンプレートマップ（Totani Fig.14 相当）Bin6 (20.76 GeV)")
    add_sep(sl)
    img = Path(f"data/figure-dwarfs/{name}/templates_fig14_{name}.png")
    if img.exists():
        sl.shapes.add_picture(str(img), Inches(0.15), Inches(0.65),
                              width=W-Inches(0.3))
    add_note(sl, "(a) 等方背景テンプレート  (b) GALPROP 銀河拡散放射  (c) 4FGL-DR2 既知点源")
    return sl


# ── スライドを全部末尾に追加してから順番を入れ替える ─────────────────────

# 挿入するスライドを生成（末尾に追加される）
new_slides_info = []
for name, label, insert_after_title in DWARFS:
    s1 = make_13bin_slide(name, label)
    s2 = make_skymaps_low_slide(name, label)
    s3 = make_skymaps_high_slide(name, label)
    s4 = make_template_slide(name, label)
    new_slides_info.append((name, insert_after_title, [s1, s2, s3, s4]))

print(f"新規スライド作成: {sum(len(v[2]) for v in new_slides_info)} 枚")

# ── 挿入位置を決定して並び替え ────────────────────────────────────────────
total = len(prs.slides)
xml_slides = prs.slides._sldIdLst
original_elems = list(xml_slides)

def find_slide_by_title(keyword):
    """タイトルにキーワードを含む最後のスライドのインデックス（0-based）"""
    result = None
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and keyword in sh.text_frame.text:
                result = i
                break
    return result

# 新スライドのインデックス（末尾から数える）
n_new = sum(len(v[2]) for v in new_slides_info)
new_start_idx = total - n_new

# 挿入後の新しい順番を計算
# まず元のスライドの順番リストを作る（新スライドなし）
order = list(range(total - n_new))  # 元のスライド

# 各天体のスライドを正しい位置に挿入
new_slide_counter = total - n_new  # 新スライドの開始インデックス

inserts = []  # (挿入位置, [新スライドインデックスリスト])
for name, insert_after_title, slides in new_slides_info:
    insert_after = find_slide_by_title(insert_after_title)
    if insert_after is not None:
        n = len(slides)
        new_idxs = list(range(new_slide_counter, new_slide_counter + n))
        inserts.append((insert_after, new_idxs))
        new_slide_counter += n
        print(f"  {name}: スライド{insert_after+1}の後に{n}枚挿入")
    else:
        print(f"  {name}: 挿入位置が見つかりません（'{insert_after_title}'）")
        new_slide_counter += len(slides)

# 挿入位置でソートして逆順に処理（後ろから挿入することで位置ズレを防ぐ）
inserts.sort(key=lambda x: x[0], reverse=True)
for insert_after, new_idxs in inserts:
    # insert_after+1 の位置に new_idxs を挿入
    for j, nidx in enumerate(reversed(new_idxs)):
        xml_slides.remove(original_elems[nidx])
        xml_slides.insert(insert_after + 1, original_elems[nidx])

print(f"\n最終スライド数: {len(prs.slides)}")
prs.save(str(PPT_PATH))
print("保存完了")
