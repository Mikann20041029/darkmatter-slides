"""
2026-07-13: 頑健性検証(領域分割再フィット)の結果を示す新規スライドを追加する。

Fig11-13相当(スライド33-35)の直後に挿入し、「なぜBin6の8.73σを本物の検出と
主張しないのか」の証拠を示す。code/diagnose_halo_degeneracy.py /
diagnose_halo_degeneracy_2.py の出力を使用。
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
PPT_PATH = REPO / "PPT" / "slides_v3_detailed_fixed2_final.pptx  -  Repaired.pptx"
FIG = REPO / "data" / "figure-halo-diagnostics" / "bin6_loglik_map.png"

prs = Presentation(str(PPT_PATH))
W = prs.slide_width
H = prs.slide_height
C_BG = RGBColor(0x05, 0x05, 0x1A)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE = RGBColor(0xFF, 0x44, 0x00)
C_CYAN = RGBColor(0x00, 0xFF, 0xFF)
C_GREEN = RGBColor(0x44, 0xFF, 0x88)
C_RED = RGBColor(0xFF, 0x44, 0x44)
C_GRAY = RGBColor(0xAA, 0xAA, 0xAA)


def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(sl):
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = C_BG


def title(sl, text, color=C_TITLE, size=22):
    tb = sl.shapes.add_textbox(Inches(0.15), Inches(0.05), W - Inches(0.3), Inches(0.55))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = True


def sep(sl):
    b = sl.shapes.add_shape(1, Inches(0.15), Inches(0.62), W - Inches(0.3), Inches(0.018))
    b.fill.solid()
    b.fill.fore_color.rgb = C_CYAN
    b.line.fill.background()


def txt(sl, lines, left, top, width, height, size=12):
    tb = sl.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in lines:
        t, c, b2, s = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(2)
        r = p.add_run()
        r.text = t
        r.font.size = Pt(s)
        r.font.color.rgb = c
        r.font.bold = b2


def make_slide():
    sl = blank()
    bg(sl)
    title(sl, "【頑健性検証】Bin5-6の見かけの検出はROI分割で符号反転する — 本物の信号ではない")
    sep(sl)

    txt(sl, [
        ("検証方法", C_CYAN, True, 13),
        ("ROIを①フェルミバブル領域(|l|<22°,10°<|b|<55°) ②それ以外の10-30° ③それ以外の30-60°"
         "の3つに分割し、それぞれの部分集合だけでMCMC同時フィットを再実行", C_WHITE, False, 11),
        ("", C_WHITE, False, 4),
        ("結果(f_halo中央値, 有意度)", C_CYAN, True, 13),
        ("Bin6(20.76GeV):  全ROI +0.67(+8.7σ)  |  ①バブル領域 +1.73(+7.8σ)  |  "
         "③バブル外・|b|≥30° −4.16(−9.0σ)", C_GREEN, True, 12),
        ("Bin5(12.29GeV):  全ROI +0.62(+8.7σ)  |  ①バブル領域 +5.29(+9.8σ)  |  "
         "③バブル外・|b|≥30° −18.8(−15.8σ)", C_GREEN, True, 12),
        ("", C_WHITE, False, 4),
        ("解釈", C_CYAN, True, 13),
        ("真のNFW ρ²ダークマターハロー信号なら、振幅は落ちても符号は反転しないはず。"
         "①と③で符号が逆転しているのは、性質の異なる2つの系統誤差が全ROI平均で"
         "偶然打ち消し合っていることを強く示唆する。", C_WHITE, False, 11),
        ("・①の正の超過 → フェルミバブルテンプレート(平坦スペクトル外挿)の過小評価の可能性", C_GRAY, False, 11),
        ("・③の負の超過 → GALPROP単一テンプレート(gas/ICS分離なし)の緯度分布不一致の可能性", C_GRAY, False, 11),
        ("(いずれも推定。検証には galprop.stanford.edu 復旧後のgas/ICS分離webrunが必要)", C_GRAY, False, 10),
        ("", C_WHITE, False, 4),
        ("結論", C_RED, True, 13),
        ("Bin6の8.73σを「Totaniの20GeV超過の再現」として主張することはできない", C_RED, True, 13),
        ("生成: code/diagnose_halo_degeneracy.py, diagnose_halo_degeneracy_2.py", C_GRAY, False, 9),
        ("出力: data/figure-halo-diagnostics/", C_GRAY, False, 9),
    ], Inches(0.2), Inches(0.72), Inches(5.6), H - Inches(0.9))

    if FIG.exists():
        sl.shapes.add_picture(str(FIG), Inches(5.95), Inches(0.75), width=Inches(6.2))
        txt(sl, [
            ("Bin6: ピクセル別log-likelihood改善量(with-halo − no-halo)。"
             "改善(赤)は|b|=10-30°・銀河中心付近に集中し、高緯度まで滑らかに広がっていない", C_GRAY, False, 9),
        ], Inches(5.95), Inches(6.55), Inches(6.2), Inches(0.6))

    return sl


new_slide = make_slide()

# スライド35(Fig13相当)の直後に挿入
insert_after = None
for i, sl in enumerate(prs.slides):
    for sh in sl.shapes:
        if sh.has_text_frame and "Totani Fig.13相当" in sh.text_frame.text:
            insert_after = i
            break
    if insert_after is not None:
        break
assert insert_after is not None, "slide35(Fig13相当)が見つからない"

xml_slides = prs.slides._sldIdLst
elems = list(xml_slides)
new_elem = elems[-1]  # 直前にadd_slideした新スライド
xml_slides.remove(new_elem)
xml_slides.insert(insert_after + 1, new_elem)

print(f"挿入位置: 旧スライド{insert_after+1}の直後 (新スライド番号は{insert_after+2})")
print(f"スライド総数: {len(prs.slides)}")
prs.save(str(PPT_PATH))
print("保存完了")
