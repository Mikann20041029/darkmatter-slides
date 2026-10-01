"""
新スライド（既存35枚以降）のフォントサイズを既存デザインに合わせる。

変更内容:
  13pt → 14pt  (本文テキスト)
  15pt → 16pt  (▼/▶ で始まるセクションヘッダ)

既存スライド(1-35)には一切手を触れない。
図スライド（【再現図】【補足図】）も対象外。
"""

from pptx import Presentation
from pptx.util import Pt

PPTX_PATH = "/mnt/c/Users/arsei/Downloads/slides_v3_detailed.pptx"
EXISTING_COUNT = 35  # 既存スライド枚数（0-indexed: 0-34 を保護）

# 図スライドキーワード（テキスト修正しない）
FIGURE_KEYWORDS = ["【再現図】", "【補足図】"]


def get_slide_title(slide) -> str:
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return shape.text_frame.text.strip()
    return ""


def is_header_para(para) -> bool:
    text = "".join(r.text for r in para.runs).strip()
    return text.startswith("▼") or text.startswith("▶")


def fix_slide_fonts(slide) -> tuple[int, int]:
    count_13 = 0
    count_15 = 0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            header = is_header_para(para)
            for run in para.runs:
                if run.font.size is None:
                    continue
                pt = round(run.font.size / 12700)
                if pt == 13:
                    run.font.size = Pt(14)
                    count_13 += 1
                elif pt == 15 and header:
                    run.font.size = Pt(16)
                    count_15 += 1
    return count_13, count_15


def main():
    prs = Presentation(PPTX_PATH)
    slides = list(prs.slides)
    print(f"総スライド数: {len(slides)}  (既存 {EXISTING_COUNT} 枚は保護)")

    total_13 = total_15 = 0
    for i, slide in enumerate(slides):
        if i < EXISTING_COUNT:
            continue
        title = get_slide_title(slide)
        if any(kw in title for kw in FIGURE_KEYWORDS):
            print(f"  [SKIP] Slide {i+1}: 図スライド")
            continue
        c13, c15 = fix_slide_fonts(slide)
        if c13 + c15 > 0:
            print(f"  [FIX]  Slide {i+1}: 13→14: {c13}件, 15→16(ヘッダ): {c15}件  {title[:40]!r}")
        else:
            print(f"  [OK]   Slide {i+1}: 変更なし  {title[:40]!r}")
        total_13 += c13
        total_15 += c15

    prs.save(PPTX_PATH)
    print(f"\n完了: 13→14pt: {total_13}件, 15→16pt: {total_15}件")
    print(f"保存先: {PPTX_PATH}")


if __name__ == "__main__":
    main()
