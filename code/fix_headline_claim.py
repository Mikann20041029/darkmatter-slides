"""
2026-07-13 ヘッドライン主張の訂正。

diagnose_halo_degeneracy.py / diagnose_halo_degeneracy_2.py での頑健性検証により、
Bin6(20.76GeV)の「f_halo=+0.665, 8.73σ」という見かけの超過は、ROIをバブル領域(|l|<22°,
10°<|b|<55°)とそれ以外の高緯度域(|b|>=30°, バブル領域除外)に分割すると符号が反転する
ことが判明した(バブル領域内: f_halo=+1.73, 7.78σ / バブル外高緯度: f_halo=-4.16, 9.04σ)。
Bin5(12.29GeV)でも同じパターン(+5.29,9.78σ / -18.8,15.8σ)。

全ROIでの「検出」は、符号・大きさの異なる2つの強い系統誤差(フェルミバブルの平坦スペクトル
外挿の不正確さ、GALPROP単一テンプレートのgas/ICS分離なしによる緯度分布の不一致)が
偶然打ち消し合って生じた見かけの結果である可能性が高いと判断し、PPT全体の
ヘッドライン主張(「Totaniの20GeV超過を再現」「8.73σ検出」)を撤回し、
頑健性検証の結果を主結果として提示する形に書き換える。

対象: スライド2(Abstract), 5(Introduction目標), 21(Results冒頭), 29(全13ビンスペクトル
キャプション), 80(Conclusion), 104(MCMC本体スライド)
"""
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
PPT_PATH = ROOT / "PPT" / "slides_v3_detailed_fixed2_final.pptx  -  Repaired.pptx"


def by_text(shapes, needle):
    for sh in shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise ValueError(f"shape not found containing {needle!r}")


def set_single_run(shape, text):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def main():
    prs = Presentation(str(PPT_PATH))

    # ---------------- slide2: Abstract ----------------
    sl = prs.slides[1]
    sh = by_text(sl.shapes, "31.0σ の有意な過剰を検出")
    set_single_run(sh, (
        "・Fermi-LAT 780週（15年）の全データでTotani (2025) の解析パイプラインを再現実装\n"
        "・GALPROP直接フィット・Loop I物理2シェルモデル・Poisson尤度MCMC同時フィットを構築\n"
        "・Bin6（20.76 GeV）でf_halo=+0.67（8.73σ）の見かけの超過を得たが、"
        "頑健性検証（ROIをフェルミバブル領域と高緯度域に分割した再フィット）で符号反転\n"
        "  → バブル領域内は+7.8σ、バブル領域外の高緯度(|b|≥30°)は−9.0σと、"
        "符号・大きさの異なる2つの系統誤差が打ち消し合った見かけの結果と判断\n"
        "  → 現時点でTotaniの20GeV超過の頑健な再現は主張できない\n"
        "・矮小銀河5天体（Draco, Sculptorなど）は全て非検出（Coma Ber.の見かけの超過も点源汚染と判明）"
    ))

    # ---------------- slide5: Introduction目標 ----------------
    sl = prs.slides[4]
    sh = by_text(sl.shapes, "結果：Bin6（20.76")
    set_single_run(sh, "結果：Bin6（20.76 GeV）で見かけ8.73σ（頑健性検証で符号反転、次項参照）")

    # ---------------- slide21: Results冒頭 ----------------
    sl = prs.slides[20]
    sh = by_text(sl.shapes, "本図は旧OLS手法時代のもの")
    set_single_run(sh, (
        "（本図は旧OLS手法時代のもの。現行MCMC結果はスライド29参照。"
        "ただしスライド33b「頑健性検証」の通り、現行のBin6検出も符号反転するため信頼できない）"
    ))

    # ---------------- slide29: 全13ビンスペクトル(headline figure) ----------------
    sl = prs.slides[28]
    sh = by_text(sl.shapes, "NFW ハロー振幅")
    p = sh.text_frame.paragraphs[0]
    p.runs[0].text = (
        "全13エネルギービンでの NFW ハロー振幅 f_halo とその1σ不確かさ（MCMC同時フィット、5自由パラメータ、等方背景はビンごとに固定）。"
    )
    p.runs[1].text = (
        "Bin 5-6 (12.3-20.8 GeV) で見かけの最大有意性(+8.7σ程度)。ただし頑健性検証で符号反転しており本物の検出とは断定できない（次項スライド参照）。"
    )
    notes_tf = sl.notes_slide.notes_text_frame
    notes_tf.text += (
        "\n\n【2026-07-13 頑健性検証の結果、重要】\n"
        "Bin5-6の見かけの超過は、ROIをフェルミバブル領域(|l|<22°,10°<|b|<55°)と"
        "それ以外の高緯度域(|b|>=30°)に分割すると符号が反転する(Bin6: バブル領域内+7.78σ, "
        "バブル外高緯度-9.04σ)。全ROIでの見かけの検出は、この2つの逆符号の系統誤差が"
        "打ち消し合った結果である可能性が高い。詳細は次項スライドと"
        "data/figure-halo-diagnostics/region_partition_summary.txt参照。"
    )

    # ---------------- slide80: Conclusion ----------------
    sl = prs.slides[79]
    shapes = sl.shapes
    by_text(shapes, "パイプライン実装完了").text_frame.paragraphs[0].runs[0].text = \
        "✓ パイプライン実装・頑健性検証を完了"
    by_text(shapes, "6成分(等方+GALPROP+LoopI").text_frame.paragraphs[0].runs[0].text = \
        "Fermi-LAT 780週・6成分(等方+GALPROP+LoopI×2+バブル+ハロー)MCMC同時フィットを実装し、頑健性を検証"
    by_text(shapes, "8.73σ を検出").text_frame.paragraphs[0].runs[0].text = \
        "✗ Bin6（20.76 GeV）の見かけの8.73σは頑健性検証で符号反転"
    sh = by_text(shapes, "13–19σ")
    p = sh.text_frame.paragraphs[0]
    p.runs[0].text = (
        "ROIをフェルミバブル領域と高緯度域に分割すると、バブル領域内+7.8σ・バブル外高緯度−9.0σと符号が逆転。"
    )
    p.runs[1].text = "2つの系統誤差(バブル外挿・GALPROP単一テンプレート)が打ち消し合った見かけの結果と判断し、Totani再現は主張しない"
    by_text(shapes, "矮小銀河5天体展開済み").text_frame.paragraphs[0].runs[0].text = \
        "✓ 矮小銀河5天体展開済み（全て非検出）"
    by_text(shapes, "GALPROP gas/ICS分離実装（Stanford webrun復旧待ち）・系統誤差の定量評価・卒論執筆").text_frame.paragraphs[0].runs[0].text = \
        "GALPROP gas/ICS分離実装・フェルミバブルスペクトル形状の精緻化（Stanford webrun復旧待ち）。"\
        "現状は「Totaniの20GeV超過を頑健に再現できなかった」ことと、その原因の特定(バブル外挿・GALPROP系統誤差)が本研究の結論"

    # ---------------- slide104: MCMC本体スライド ----------------
    sl = prs.slides[103]
    sh = by_text(sl.shapes, "MCMC同時フィット")
    sh.text_frame.paragraphs[0].runs[0].text = \
        "MCMC同時フィット — 本研究の本体手法（頑健性検証で見かけの検出と判明、2026-07-13）"
    sh2 = by_text(sl.shapes, "Bin6(20.76GeV)現行値")
    sh2.text_frame.paragraphs[0].runs[0].text = (
        "Bin6(20.76GeV)現行値: f_gal=0.69, f_halo=0.67, 見かけの有意度8.7σ。"
        "ただしROIをバブル領域/高緯度域に分割すると符号反転(+7.8σ/−9.0σ)することを確認済み"
    )
    sh3 = by_text(sl.shapes, "全13ビンの現行結果は")
    sh3.text_frame.paragraphs[0].runs[0].text = (
        "全13ビンの現行結果は results/mcmc_allbins/halo_spectrum.json 参照。"
        "頑健性検証(領域分割再フィット)の詳細は data/figure-halo-diagnostics/ 参照。"
        "現時点でTotaniの20GeV超過の頑健な再現は主張できない"
    )

    prs.save(str(PPT_PATH))
    print("保存完了:", PPT_PATH)


if __name__ == "__main__":
    main()
