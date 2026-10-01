"""
issue-1 (中間発表PPT107スライド精査) 🔴優先度スライドの本文更新。

対象 (旧OLS/exp近似時代の記述が残存、または現行MCMC結果と不整合だったスライド):
  slide29  【再現図】全13ビンNFW振幅スペクトル — 軽微な事実誤り修正(6→5自由パラメータ, OLS表記除去)
  slide40  Discussion(2/3) 系統誤差 — 統計手法/露出マップは解消済み、GALPROP gas/ICS分離が現在の最大の残存差と明記
  slide80  Conclusion — 旧OLS 31.03σ → 現行MCMC 8.73σ、Coma Ber. 4.54σ→4.04σ(778/780週本番)に修正
  slide92  6成分定量比較 — Loop I幾何モデル一致・点源DR4一致・NFW有意度をMCMC現行値に更新
  slide93  全13ビンスペクトル表 — 旧all_bins_nfw_fit.py(OLS)値→results/mcmc_allbins/halo_spectrum.json値に置換
  slide104 画像を旧Bin6単体テスト時のcornerプロットから現行13ビンスペクトル(halo_spectrum.png)に差し替え

.dev/issues/issue-1/TODO.md の T-1.029(既完了・軽微訂正のみ)/T-1.040/T-1.080/T-1.092/T-1.093/T-1.104 に対応。
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
PPT_PATH = ROOT / "PPT" / "slides_v3_detailed_fixed2_final.pptx  -  Repaired.pptx"
HALO_SPECTRUM_PNG = ROOT / "results" / "mcmc_allbins" / "halo_spectrum.png"


def set_para(paragraph, *texts):
    """paragraph の runs[0]に texts[0] を、runs[1]に texts[1] を…と割り当てる。
    run 数が texts より多い場合は残りを空文字にする。run が texts より少ない場合はエラー
    (このスクリプトでは事前に run 数を確認済みの箇所のみ呼ぶ)。"""
    runs = paragraph.runs
    assert len(texts) <= len(runs), f"texts({len(texts)}) > runs({len(runs)}): {[r.text for r in runs]}"
    for i, r in enumerate(runs):
        r.text = texts[i] if i < len(texts) else ""


def by_text(shapes, needle):
    for sh in shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise ValueError(f"shape not found containing {needle!r}")


def main():
    prs = Presentation(str(PPT_PATH))
    slides = prs.slides

    # ------------------------------------------------------------------
    # slide29 (0-idx 28): 「6自由パラメータ」→ 実際は5(iso固定)。「OLS過大評価を含む」は
    # 現行版がMCMCそのものになったため事実と矛盾するので除去。
    # ------------------------------------------------------------------
    sl = slides[28]
    sh = by_text(sl.shapes, "NFW ハロー振幅")
    p = sh.text_frame.paragraphs[0]
    set_para(
        p,
        "全13エネルギービンでの NFW ハロー振幅 f_halo とその1σ不確かさ（MCMC同時フィット、5自由パラメータ、等方背景はビンごとに固定）。",
        "Bin 5-6 (12.3-20.8 GeV) で最大の有意性(+8.7σ程度)を検出。",
    )

    # ------------------------------------------------------------------
    # slide40 (0-idx 39): Discussion(2/3) 系統誤差テーブル
    # ------------------------------------------------------------------
    sl = slides[39]
    shapes = sl.shapes

    # 行1: 統計手法 — 現行は本研究・TotaniともにMCMC+Poisson尤度で一致(解消)
    by_text(shapes, "最小二乗法").text_frame.paragraphs[0].runs[0].text = \
        "MCMC + Poisson尤度\n(2026-07以降)"
    by_text(shapes, "MCMC + Poisson尤度\n").text_frame.paragraphs[0].runs[0].text = \
        "MCMC + Poisson尤度"
    by_text(shapes, "大").text_frame.paragraphs[0].runs[0].text = "解消"
    sh = by_text(shapes, "S/N が過大評価")
    sh.text_frame.paragraphs[0].runs[0].text = (
        "統計手法は両者ともMCMC+Poisson尤度で一致(2026-07-12までに移行済み)\n"
        "残る差はBin6有意度(本研究8.73σ vs Totani報告13-19σ)。主因はGALPROP gas/ICS分離未実装(次項)"
    )

    # 行2: GALPROPモデル — 直接フィットに移行済みだが、gas/ICS分離は依然未実装(現在の最大の残存差)
    by_text(shapes, "gll_iem_v07の\n1種のみ").text_frame.paragraphs[0].runs[0].text = \
        "gll_iem_v07.fits\n直接フィット(単一合成)"
    by_text(shapes, "複数モデル比較").text_frame.paragraphs[0].runs[0].text = \
        "自己webrun(SLZ6R30\nT150C2)でgas/ICS\n独立2テンプレート"
    # 「中」は既に②行で使用済みラベルなので、直前の1つ目「解消」変更後に残る「中」を探す
    med_impact = None
    for sh_ in shapes:
        if sh_.has_text_frame and sh_.text_frame.text.strip() == "中":
            med_impact = sh_
            break
    assert med_impact is not None
    med_impact.text_frame.paragraphs[0].runs[0].text = "大"
    sh = by_text(shapes, "モデルが実際の銀河拡散放射と違う場合")
    sh.text_frame.paragraphs[0].runs[0].text = \
        "gas/ICS分離がないため現在の系統誤差の主因。galprop.stanford.edu復旧待ち(.dev/TODO.md参照)、webrun再現までGALPROP起因の系統誤差は未定量化"

    # 行3: 露出マップ — 2026-07-11に実測13ビン分を導入済み(解消)
    by_text(shapes, "未使用").text_frame.paragraphs[0].runs[0].text = "実測13ビン分導入済み\n(2026-07-11)"
    # 残る「中」(行3の影響列)を解消に変更
    med_impact2 = None
    for sh_ in shapes:
        if sh_.has_text_frame and sh_.text_frame.text.strip() == "中":
            med_impact2 = sh_
            break
    assert med_impact2 is not None
    med_impact2.text_frame.paragraphs[0].runs[0].text = "解消"
    by_text(shapes, "観測方向ごとの有効観測時間の差を補正できていない").text_frame.paragraphs[0].runs[0].text = \
        "2026-07-11にHEASARC gtexposure由来の実測露出マップ(13ビン)を導入し解消。旧仮定数(Mrk501較正1.24e11)は実測平均(5.80e11 cm²・s)の1/4.68だった"

    # 行4: 点源処理 — 変更なし(小のまま)、DR4反映のみ追記
    by_text(shapes, "座標ピクセルをNaN").text_frame.paragraphs[0].runs[0].text = \
        "4FGL-DR4座標\nピクセルをNaN化"

    # 末尾サマリ行
    by_text(shapes, "卒論での記述").text_frame.paragraphs[0].runs[0].text = \
        "→ 卒論での記述: 「統計手法・露出マップはTotaniと同一方式に移行済み(2026-07)。残る最大の差はGALPROP gas/ICS分離未実装（Stanford webrun復旧待ち）。これが現在の系統誤差の主因である。」"

    # ------------------------------------------------------------------
    # slide80 (0-idx 79): Conclusion
    # ------------------------------------------------------------------
    sl = slides[79]
    shapes = sl.shapes
    by_text(shapes, "5成分差し引き・NFWフィットを実装").text_frame.paragraphs[0].runs[0].text = \
        "Fermi-LAT 780週・6成分(等方+GALPROP+LoopI×2+バブル+ハロー)MCMC同時フィットを実装"
    by_text(shapes, "31.03σ").text_frame.paragraphs[0].runs[0].text = \
        "✓ Bin6（20.76 GeV）で 8.73σ を検出（MCMC同時フィット）"
    sh = by_text(shapes, "13–19σ")
    set_para(
        sh.text_frame.paragraphs[0],
        "Totani (2025) の 13–19σ と同オーダー。",
        "低E側(Bin1-4)には有意な負のハロー(欠損、最大22.7σ)も残る",
    )
    sh = by_text(shapes, "Coma Ber.の4.54")
    set_para(
        sh.text_frame.paragraphs[0],
        "→ ",
        "矮小銀河5天体は全て非検出（Coma Ber.の4.04",
        "σ(生, 778/780週本番データ)は点源汚染、マスク後-0.93σ）",
    )
    by_text(shapes, "GALPROP複数モデル比較・露出マップ導入による精度向上").text_frame.paragraphs[0].runs[0].text = \
        "GALPROP gas/ICS分離実装（Stanford webrun復旧待ち）・系統誤差の定量評価・卒論執筆"

    # ------------------------------------------------------------------
    # slide92 (0-idx 91): 6成分定量比較
    # ------------------------------------------------------------------
    sl = slides[91]
    sh = by_text(sl.shapes, "GALPROP 銀河拡散成分のズレ")
    tf = sh.text_frame
    set_para(tf.paragraphs[2], "▼ GALPROP 銀河拡散成分 — 2026-07-10以前の指数近似時代の偏差（歴史的参考値、現在は解消）")
    set_para(tf.paragraphs[8],
             "→ 2026-07-11に gll_iem_v07.fits の直接フィット(単一合成テンプレート)へ移行しこの偏差は解消。"
             "残る差は gas/ICS 分離(Totaniのwebrun手法)未実装")
    set_para(tf.paragraphs[10], "▼ 6成分の現状比較（2026-07-13時点）")
    set_para(tf.paragraphs[13],
             "点源マスク   4FGL-DR4 (14年)            4FGL-DR4 (14年)      カタログ世代一致（2026-07-10 DR2→DR4更新済み）")
    set_para(tf.paragraphs[14],
             "フェルミバブル 幾何残差テンプレート(Bin3基準)     正負2テンプレート+σ=1°     テンプレート構成は簡略化(正のみ・smoothingなし)。")
    tf.paragraphs[14].add_run() if False else None
    set_para(tf.paragraphs[15],
             "ループI      2シェル物理モデル(Ackermann+2014/2017)  同左（同一幾何モデル）    2026-07-12に幾何モデルを完全一致化。残る差は振幅フィット方式のみ")
    set_para(tf.paragraphs[16],
             "NFW J-factor MCMC同時フィット(emcee,Poisson尤度)  同左(MCMC同時フィット)   手法は一致。有意度はBin6=8.73σ(本研究)vsTotani報告13-19σ、主因はGALPROP gas/ICS分離未実装")

    # ------------------------------------------------------------------
    # slide93 (0-idx 92): 全13ビンスペクトル表 — MCMC現行値に置換
    # ------------------------------------------------------------------
    import json
    halo = json.load(open(ROOT / "results" / "mcmc_allbins" / "halo_spectrum.json"))
    rows = []
    for b in halo["bins"]:
        fh = b["params"]["f_halo"]["median"]
        sig = b["significance_sigma"]
        signed = sig if fh >= 0 else -sig
        sign_txt = "正" if fh >= 0 else "負"
        mark = " ★" if b["bin"] in (5, 6) else ""
        rows.append(f" {b['bin']:02d}  {b['e_center_gev']:7.2f}   {signed:+6.2f} σ{mark}   {sign_txt}")

    sl = slides[92]
    sh = by_text(sl.shapes, "結果一覧")
    tf = sh.text_frame
    set_para(tf.paragraphs[1], "▼ 結果一覧（mcmc_fit_all_bins.py, 2026-07-12 Loop I幾何モデル修正+多点始動化後, commit 0128fbe）")
    set_para(tf.paragraphs[2], "Bin  E[GeV]    有意度[σ]   f_halo符号")
    for i, row in enumerate(rows):
        set_para(tf.paragraphs[3 + i], row)
    set_para(tf.paragraphs[16],
             "Bin5-6(12-21GeV)で最大+8.7σ程度、Bin1-4は最大-22.7σの負の有意度(参照: results/mcmc_allbins/halo_spectrum.json)")

    sh2 = by_text(sl.shapes, "スペクトルの読み方")
    tf2 = sh2.text_frame
    set_para(tf2.paragraphs[0], "▼ スペクトルの読み方")
    set_para(tf2.paragraphs[1], "有意度>0: NFW 形状の超過あり（DM 候補）")
    set_para(tf2.paragraphs[2], "有意度<0: 差引きモデルの過大評価（欠損）")
    set_para(tf2.paragraphs[4], "▼ 観測されたスペクトル形状（b クォーク対消滅 DM との対比）")
    set_para(tf2.paragraphs[5], "・Bin5-6 (12.3-20.8 GeV) でピーク的な正の超過(+8.7σ程度) ✓ Totaniの20GeVピークと定性的に整合")
    set_para(tf2.paragraphs[6], "・低エネルギー(<7.3 GeV)は逆に強い負の有意度(最大-22.7σ) — GALPROPガス成分不確かさに由来と推定(Discussion参照)")
    set_para(tf2.paragraphs[7], "・高エネルギー(>280 GeV)にも+2.8〜3.0σの残留超過 — 統計が乏しく系統誤差の可能性を否定できない")
    set_para(tf2.paragraphs[8], "→ 低エネルギー側の大きな負の有意度は理想的なDMスペクトル形状とは一致しない")
    set_para(tf2.paragraphs[10], "▼ 図: スライド29（Results全13ビンNFW振幅スペクトル）を参照")
    set_para(tf2.paragraphs[11], "（数値ソース: results/mcmc_allbins/halo_spectrum.json, 生成: code/mcmc_fit_all_bins.py）")

    # ------------------------------------------------------------------
    # slide104 (0-idx 103): 画像を旧Bin6単体corner plotから現行13ビンスペクトルに差し替え
    # ------------------------------------------------------------------
    sl = slides[103]
    for sh in sl.shapes:
        if sh.shape_type == 13:
            from PIL import Image
            w, h = Image.open(HALO_SPECTRUM_PNG).size
            new_h = Emu(sh.height)
            new_w = Emu(int(sh.height * w / h))
            left = sh.left + (sh.width - new_w) // 2
            sh._element.getparent().remove(sh._element)
            sl.shapes.add_picture(str(HALO_SPECTRUM_PNG), left, sh.top, width=new_w, height=new_h)
            break

    prs.save(str(PPT_PATH))
    print("保存完了:", PPT_PATH)


if __name__ == "__main__":
    main()
