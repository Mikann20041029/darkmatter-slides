"""
2026-07-13 セルフレビューで発見した記述の訂正。

fix_issue1_late_stage_slides.py で書いた「主因はGALPROP gas/ICS分離未実装」という
文言は、Bin6有意度ギャップ(本研究8.73σ vs Totani報告13-19σ)の"原因"を検証なしに
断定していた(speculationをfactとして書いてしまっていた)。実際に確認済みなのは
「4項目の既知の手法差のうち3つ(統計手法・露出マップ・Loop I幾何)は解消し、GALPROP
gas/ICS分離だけが未解消のまま残っている」という事実のみで、この差が具体的に
何σ分のギャップに相当するかは一切テストしていない。スライド40・92の該当4箇所を
「未検証である」ことが明確に伝わる表現に訂正する。
"""
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
PPT_PATH = ROOT / "PPT" / "slides_v3_detailed_fixed2_final.pptx  -  Repaired.pptx"


def by_text(shapes, needle):
    for sh in shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise ValueError(f"shape not found containing {needle!r}")


def main():
    prs = Presentation(str(PPT_PATH))

    sl = prs.slides[39]  # slide40
    sh = by_text(sl.shapes, "統計手法は両者とも")
    sh.text_frame.paragraphs[0].runs[0].text = (
        "統計手法は両者ともMCMC+Poisson尤度で一致(2026-07-12までに移行済み)\n"
        "残る差はBin6有意度(本研究8.73σ vs Totani報告13-19σ)。原因は未特定・複数要因の可能性"
        "(GALPROP gas/ICS分離未実装は既知の手法差では最大だが、ギャップへの寄与は未検証、次項)"
    )

    sh = by_text(sl.shapes, "gas/ICS分離がないため")
    sh.text_frame.paragraphs[0].runs[0].text = (
        "統計手法・露出マップの解消後に残る、本解析が識別済みの手法差としては最大のもの。"
        "galprop.stanford.edu復旧待ち(.dev/TODO.md参照)。webrun再現までこの差が有意度ギャップに"
        "与える寄与は未定量化・未検証"
    )

    sh = by_text(sl.shapes, "卒論での記述")
    sh.text_frame.paragraphs[0].runs[0].text = (
        "→ 卒論での記述: 「統計手法・露出マップはTotaniと同一方式に移行済み(2026-07)。"
        "本解析が識別済みの手法差として残る最大のものはGALPROP gas/ICS分離未実装"
        "（Stanford webrun復旧待ち）。ただしこれが8.73σ vs 13-19σの差にどれだけ寄与するかは"
        "未検証であり、断定しない。」"
    )

    sl = prs.slides[91]  # slide92
    sh = by_text(sl.shapes, "NFW J-factor MCMC同時フィット")
    tf = sh.text_frame
    for p in tf.paragraphs:
        if p.runs and "NFW J-factor MCMC同時フィット" in p.runs[0].text:
            p.runs[0].text = (
                "NFW J-factor MCMC同時フィット(emcee,Poisson尤度)  同左(MCMC同時フィット)   "
                "手法は一致。有意度はBin6=8.73σ(本研究)vsTotani報告13-19σ。既知の手法差として"
                "残るのはGALPROP gas/ICS分離未実装だが、この差がギャップの原因であるとは未検証"
            )
            break

    prs.save(str(PPT_PATH))
    print("保存完了:", PPT_PATH)


if __name__ == "__main__":
    main()
