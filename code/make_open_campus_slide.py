"""大学説明会用に「中村誠一が研究室で何をしているか」を1枚にまとめたスライドを作る。

想定読者: 高校生・受験生 (一般向け)。専門用語を避け、図1枚で研究の全体像を伝える。
出力: docs/slide_open-campus_nakamura-intro.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
FIGURE = ROOT / "data" / "figure-week780-nfw-fit" / "nfw_halo_spectrum.png"
OUT = ROOT / "docs" / "slide_open-campus_nakamura-intro.pptx"

# 既存スライド (redesign_new_slides.py) のダークテーマに合わせる
C_BG = RGBColor(0x05, 0x05, 0x1A)
C_TITLE = RGBColor(0xFF, 0xCC, 0x00)
C_NAME = RGBColor(0x44, 0xCC, 0xFF)
C_BODY = RGBColor(0xFF, 0xFF, 0xFF)
C_CAPTION = RGBColor(0xAA, 0xAA, 0xAA)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 白紙レイアウト
    set_background(slide, C_BG)

    # --- タイトル ---
    title_tf = add_text_box(slide, Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.7))
    p = title_tf.paragraphs[0]
    run = p.add_run()
    run.text = "見えない宇宙の正体を追う ― ガンマ線でダークマターを探す"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = C_TITLE
    run.font.name = "Meiryo"

    # --- 発表者名 ---
    name_tf = add_text_box(slide, Inches(0.6), Inches(1.05), Inches(12.1), Inches(0.4))
    p = name_tf.paragraphs[0]
    run = p.add_run()
    run.text = "中村誠一（理学部物理学科 4年・卒業研究生）"
    run.font.size = Pt(18)
    run.font.color.rgb = C_NAME
    run.font.name = "Meiryo"

    # --- 本文（左側の説明） ---
    body_tf = add_text_box(slide, Inches(0.6), Inches(1.7), Inches(6.5), Inches(4.6))
    bullets = [
        "宇宙にある物質の約85%は、光を出さず正体もわからない"
        "「ダークマター」",
        "NASAのガンマ線天文衛星「Fermi」が15年間集めた、"
        "銀河中心方向のデータ（780週分）を解析",
        "ダークマター同士がぶつかって消えるときに出ると予言"
        "される「20 GeV」のガンマ線の痕跡を探索",
        "右の図：エネルギーごとの信号の強さを表したグラフ。"
        "20 GeV付近（オレンジの帯）だけ周りより強い信号が"
        "現れている → ダークマターからの信号の候補",
    ]
    for i, text in enumerate(bullets):
        p = body_tf.paragraphs[0] if i == 0 else body_tf.add_paragraph()
        p.space_after = Pt(14)
        run = p.add_run()
        run.text = "・" + text
        run.font.size = Pt(16)
        run.font.color.rgb = C_BODY
        run.font.name = "Meiryo"

    # 結び（先行研究との関係）
    closing_tf = add_text_box(slide, Inches(0.6), Inches(6.5), Inches(6.5), Inches(0.8))
    p = closing_tf.paragraphs[0]
    run = p.add_run()
    run.text = "→ 先行研究（Totani, 2025）が報告した信号を、自分の手で独立に再現・検証している"
    run.font.size = Pt(14)
    run.font.italic = True
    run.font.color.rgb = C_NAME
    run.font.name = "Meiryo"

    # --- 図（右側） ---
    img_w = Inches(5.6)
    img_left = Inches(7.4)
    img_top = Inches(1.7)
    pic = slide.shapes.add_picture(str(FIGURE), img_left, img_top, width=img_w)

    # キャプション（図の下）
    cap_top = Emu(img_top + pic.height + Inches(0.15))
    cap_tf = add_text_box(slide, img_left, cap_top, img_w, Inches(1.2))
    cap_tf.word_wrap = True
    p = cap_tf.paragraphs[0]
    run = p.add_run()
    run.text = (
        "図: ガンマ線の強さをエネルギーごとに並べたグラフ。"
        "横軸がエネルギー、縦軸が信号の強さ。"
        "オレンジの帯（20 GeV付近）だけ突出して高い値を示しており、"
        "ダークマター同士の対消滅で生じる信号と矛盾しない。"
    )
    run.font.size = Pt(12)
    run.font.color.rgb = C_CAPTION
    run.font.name = "Meiryo"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
