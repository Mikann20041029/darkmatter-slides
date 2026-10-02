"""_claude 版の「手法①：テンプレートフィット」の後に、各成分 (型紙) がどうやってガンマ線を出すかの説明を 7 枚足す。
右側には各成分の 20.8 GeV の型紙の地図を置く (cloud_reports/2026-10-02_plot_templates.py で作成)。
実行: python slides/figs/add_component_slides.py <入力 pptx> <出力 pptx>
"""
import copy
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
L = {l.name: l for l in prs.slide_layouts}
slides = list(prs.slides)
REF = next(sh for s in slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Reference: D. Foreman"))
anchor = next(i for i, s in enumerate(slides) if s.shapes.title is not None and s.shapes.title.text.startswith("手法①"))
TD = "slides/figs/templates/template_"
R_T = "T. Totani (2025), arXiv:2507.07209 §2.3"


def slide(title, items, pics, note, ref=R_T, size=18):
    s = prs.slides.add_slide(L["Title and Content"])
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(30)
    b = s.placeholders[1]
    b.left, b.top, b.width, b.height = Inches(0.6), Inches(1.5), Inches(7.0), Inches(4.7)
    tf = b.text_frame
    for k, it in enumerate(items):
        lvl, t = (1, it[1]) if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.text, p.level = t, lvl
        for r in p.runs:
            r.font.size = Pt(size - 3 if lvl else size)
    if len(pics) == 1:
        s.shapes.add_picture(TD + pics[0] + ".png", Inches(7.9), Inches(1.3), height=Inches(4.9))
    else:
        for j, pname in enumerate(pics):
            s.shapes.add_picture(TD + pname + ".png", Inches(7.75 + 2.75 * j), Inches(2.0), width=Inches(2.7))
    el = copy.deepcopy(REF._element)
    s.shapes._spTree.append(el)
    p = s.shapes[-1].text_frame.paragraphs[0]
    p.runs[1].text = " " + ref
    for r in p.runs[2:]:
        r._r.getparent().remove(r._r)
    s.notes_slide.notes_text_frame.text = note + "\n右の地図: 20.8 GeV の型紙 (倍率は掛けていない。形だけを見る図)。cloud_reports/2026-10-02_plot_templates.py で作成。"
    return s


new = [
    slide("成分①：等方背景", [
        "空のどの方向からも、同じ明るさで来るガンマ線",
        "出どころ：天の川の外の遠い銀河（ブレーザーなど）",
        ("", "一つ一つは見分けられないほど暗い光が、たくさん重なったもの"),
        "型紙：空全体で一様（どこも同じ明るさ）",
        "倍率：0 以上",
        ("", "最初の値は E²dN/dE = 10⁻⁴ MeV cm⁻² s⁻¹ sr⁻¹ 相当"),
        "注意：Loop I の殻 1 と形がほぼ同じで、区別しにくい",
    ], ["iso_counts"],
       "たとえ: 部屋の壁全体がぼんやり同じ明るさで光っている、その「底上げ」の明るさ。\n"
       "最初の値 10⁻⁴ は、これまでに測られた「銀河の外から来るガンマ線の背景」の典型的な明るさ (Totani §2.3)。\n"
       "Loop I との区別: 殻 1 は太陽が殻の壁の中にあるため空のほぼ全体を覆い、明るさの変化が等方背景と同じくらい平ら。データからは 2 つの分け方が決まらず、合計だけが決まる (.dev/TOTANI_SPEC.md 3.7a)。"),

    slide("成分②：ガス（π⁰ 崩壊＋制動放射）", [
        "宇宙線が、星と星の間のガス（主に水素）にぶつかって出すガンマ線",
        "π⁰ 崩壊",
        ("", "宇宙線の陽子が水素の原子核にぶつかる → π⁰ という粒子ができる → すぐ光子 2 個に壊れる"),
        "制動放射",
        ("", "宇宙線の電子がガスのそばを通るときに曲げられ、ブレーキがかかって光を出す"),
        "型紙：GALPROP で計算。2 つの比は GALPROP の値に固定して 1 枚",
        ("", "ガスの多い銀河面の近くと、ガス雲の模様が明るい"),
        "倍率：0 以上、最初は 1",
    ], ["gas"],
       "GALPROP: 宇宙線が銀河の中でどう生まれ、どう広がるかを計算し、そこから出るガンマ線の地図を作るシミュレーション (ver. 54、設定 SLZ6R30T150C2。Fermi-LAT チームの銀河拡散モデルでも使われている設定)。\n"
       "π⁰ は電気を持たない中間子で、できるとすぐ (10⁻¹⁶ 秒ほど) 光子 2 個に壊れる。\n"
       "ガス雲の模様は、電波 (水素の 21 cm 線) や、ちりによる星の光の赤み (E(B−V)) から作ったガスの地図に由来する。",
       ref=R_T + "; GALPROP v54"),

    slide("成分③：ICS（逆コンプトン散乱）", [
        "宇宙線の電子が、まわりの光にぶつかってガンマ線に変える",
        ("", "ぶつかる光：星の光、ちりが出す赤外線、宇宙背景放射"),
        "たとえ：速いボール（電子）が軽いピンポン玉（光）をはね飛ばす",
        ("", "はね飛ばされたピンポン玉がガンマ線になる"),
        "星の光は銀河の中心ほど多い",
        ("", "→ 形は、銀河面と中心付近が明るく、なめらかに広がる"),
        "型紙：GALPROP で計算",
        "倍率：0 以上、最初は 1",
        "ハローと形が似ている（縮退の原因）",
    ], ["ics"],
       "逆コンプトン散乱: 普通のコンプトン散乱は光が電子をはね飛ばす。ここでは逆で、とても速い電子が、低いエネルギーの光にエネルギーを渡して、ガンマ線にする。\n"
       "ガスの型紙と違って雲の模様が無く、なめらか。ハローもなめらかで中心に向かって明るいので、2 つの形が似ている (問題のスライドで説明する縮退)。\n"
       "GALPROP の設定は IC_anisotropic = 0 (等方近似)。",
       ref=R_T + "; GALPROP v54"),

    slide("成分④：点源", [
        "一つ一つ位置が分かっている天体",
        ("", "パルサー（回転する中性子星）、ブレーザー（遠い銀河の中心）など"),
        "型紙：Fermi-LAT の天体カタログ（4FGL-DR4）から作る",
        ("", "位置と、カタログに書かれた明るさのエネルギー変化"),
        ("", "望遠鏡のぼやけ（PSF）で点を広げる。ぼやけはエネルギーで変わる"),
        "全部の点源をまとめて 1 枚、倍率 1 個",
        "倍率：0 以上、最初は 1",
        "大きく広がった天体（拡張源）は解析から外す",
        ("", "長半径の 2 倍の円（地図の白い丸）"),
    ], ["ps"],
       "4FGL-DR4: Fermi-LAT の 14 年分のデータで見つかった天体のカタログ (gll_psc_v35.fit)。\n"
       "PSF (点広がり関数): 望遠鏡で点を撮っても少しぼやけて写る、そのぼやけ方。エネルギーが低いほど大きくぼやける (1 GeV で約 0.85°、10 GeV で約 0.16°)。\n"
       "拡張源: 大マゼラン雲や Cen A のローブのように、点ではなく広がって見える天体。Totani §2.3 と同じく、長半径の 2 倍の円を外す。"),

    slide("成分⑤：Loop I", [
        "電波で見える、空の 100° 以上に広がる巨大な輪っか",
        ("", "ガンマ線でも見えている"),
        "正体は未確定",
        ("", "近くの超新星の残骸か、銀河中心から吹き出したものか"),
        "型紙：近くの超新星の残骸とみなした形",
        ("", "半径 50〜100 pc の殻 2 つが、一様に光る"),
        ("", "殻の位置と大きさは Fermi-LAT チームの値"),
        "2 つの殻に、別々の倍率（0 以上、最初は 0）",
    ], ["loopI_a", "loopI_b"],
       "pc (パーセク) は距離の単位で、1 pc ≈ 3.26 光年。太陽から殻の中心までは 78 pc と 95 pc。\n"
       "殻 1 は太陽が殻の壁の中にあるので、空のほぼ全体を覆う (左の地図)。そのため等方背景と区別しにくい。\n"
       "超新星の残骸: 星が爆発したときに吹き飛んだガスが、泡のように広がったもの。",
       ref=R_T + " (殻の値は Fermi-LAT チームの論文)"),

    slide("成分⑥：フェルミバブル（正・負）", [
        "銀河中心の上下に、それぞれ約 50° 広がる巨大な泡",
        ("", "過去の銀河中心の活動（ブラックホールや爆発的な星形成）の名残と考えられている"),
        "型紙はデータから作る（Totani §3.1）",
        ("", "まず平らな泡の形で当てはめ、残った光（4.3 GeV）を 1° でぼかす"),
        ("", "残りが正の所 → 正の型紙（バブル本体）"),
        ("", "残りが負の所 → 負の型紙（GALPROP とデータのずれの補正）"),
        "倍率：正は 0 以上、負は符号自由。最初は 0",
    ], ["fb", "fb_neg"],
       "なぜデータから作るか: バブルの正確な形を予言する理論が無いため。Totani は Fermi-LAT チームのやり方にならい、当てはめの残りから形を作った。\n"
       "4.3 GeV の地図を使うのは、バブルがはっきり見えるため (Totani §3.1)。形はエネルギーで変えず、明るさ (倍率) だけを各エネルギーで決める。\n"
       "負の型紙: 当てはめで光が余らなかった (予測が多すぎた) 場所。主に GALPROP の予測とデータのずれ。符号自由の倍率を掛けて補正に使う。\n"
       "左の地図が正、右が負 (負の残差の大きさ)。",
       ref="T. Totani (2025), arXiv:2507.07209 §3.1"),

    slide("成分⑦：ハロー（探しているもの）", [
        "暗黒物質の粒子どうしがぶつかって消え（対消滅）、ガンマ線を出す",
        "明るさ ∝ 密度の 2 乗（粒子が 2 つ要るため）",
        ("", "それを視線方向に足し合わせたもの = J ファクター"),
        "密度の形：NFW（Via Lactea II のシミュレーション）",
        ("", "rs = 21 kpc、太陽の位置で 0.42 GeV/cm³"),
        "形：銀河中心に向かって明るく、なめらかに広がる",
        "倍率：符号自由（負も許す）、最初は 0",
        ("", "あるか分からないものを探すので、偏りを作らないため"),
    ], ["halo"],
       "NFW: シミュレーションで得られる暗黒物質ハローの典型的な密度の形。中心付近では距離に反比例、外側では距離の 3 乗に反比例して薄くなる。\n"
       "形はエネルギーで変わらない (どのエネルギーでも同じ地図)。明るさ (倍率) だけを各エネルギーで決める。だから「どのエネルギーでハローが必要か」が有意度のスペクトルになる。\n"
       "ICS と同じく、なめらかで中心に向かって明るい。違いは、ハローのほうが中心に強く集まっていること。"),
]

lst = prs.slides._sldIdLst
ids = list(lst)
for k, sid in enumerate(ids[-len(new):]):
    lst.remove(sid)
    lst.insert(anchor + 1 + k, sid)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "枚。成分の説明は", anchor + 2, "〜", anchor + 1 + len(new), "枚目")
