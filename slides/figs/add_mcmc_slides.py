"""本人の編集版に、MCMC の数式つき解説スライドを 3 枚足して _claude 版を作る。
本人のスライドは 1 枚も変えない (7 枚目「手法④：MCMC」の直後に挿入するだけ)。

実行 (リポジトリの一番上で):
  git show "origin/main:slides/2026-10-02_研究進捗報告.pptx" > /tmp/user.pptx
  python slides/figs/add_mcmc_slides.py /tmp/user.pptx "slides/2026-10-02_研究進捗報告_claude.pptx"
"""
import copy
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
L = {l.name: l for l in prs.slide_layouts}
anchor = next(i for i, s in enumerate(prs.slides) if s.shapes.title is not None and "MCMC" in s.shapes.title.text)
REF = next(sh for sh in prs.slides[anchor].shapes if sh.has_text_frame and sh.text_frame.text.startswith("Reference"))


def add_ref(s, text):
    el = copy.deepcopy(REF._element)
    s.shapes._spTree.append(el)
    p = s.shapes[-1].text_frame.paragraphs[0]
    p.runs[1].text = " " + text
    for r in p.runs[2:]:
        r._r.getparent().remove(r._r)


def slide(title, items, note, ref=None, size=22, body_w=11.5):
    s = prs.slides.add_slide(L["Title and Content"])
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(34)
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Inches(0.92), Inches(1.6), Inches(body_w), Inches(4.5)
    tf = body.text_frame
    for i, it in enumerate(items):
        lvl, txt = (1, it[1]) if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = txt, lvl
        for r in p.runs:
            r.font.size = Pt(size - 4 if lvl else size)
    if ref:
        add_ref(s, ref)
    s.notes_slide.notes_text_frame.text = note
    return s


new = []
new.append(slide("MCMC の中身①：何を求めているか", [
    "知りたいもの：データ D を見たあと、倍率 θ = (f₁, …, f₉) がどの値をとりうるか",
    "ベイズの定理：P(θ | D) = L(D | θ) π(θ) / P(D)",
    ("", "P(θ | D)：事後分布（データを見た後の、θ のもっともらしさ）"),
    ("", "L(D | θ)：尤度（手法② の L）　π(θ)：事前分布（今回は一様）"),
    ("", "P(D)：θ によらない定数 → 計算しなくてよい"),
    "つまり P(θ | D) ∝ L(D | θ) × π(θ)",
    "ただし θ は 9 次元。1 次元を 100 点に区切っても 100⁹ = 10¹⁸ 通り → 全部は計算できない",
    "→ 事後分布から「サンプルを取り出す」方法が MCMC",
], "ベイズの定理: 「データを見る前の考え (事前分布)」を「データの当てはまり (尤度)」で更新すると「データを見た後の考え (事後分布)」になる。\n"
   "事前分布は一様: 倍率の符号の制約 (ハローとバブル負以外は 0 以上) と、|f| ≤ 10⁶ の広い制限だけ。だから事後分布の形は尤度の形とほぼ同じ。\n"
   "P(D) (証拠) は全部の θ について足し合わせた値で、計算がとても大変。でも MCMC では「2 つの点の比」しか使わないので、割り算で消えて計算しなくてよい。\n"
   "たとえ: 9 次元の山の地形図を全部描くのは無理なので、山の上を歩き回って「よく通った場所 = 高い場所」を記録する。",
   ref="D. Foreman-Mackey et al. (2013), PASP 125, 306"))

s = slide("MCMC の中身②：歩き方（メトロポリス法）", [
    "① 今の点 θ から、少しずらした候補 θ′ を作る",
    "② 比 r = P(θ′ | D) / P(θ | D) を計算",
    ("", "P(D) は分母分子で消える"),
    "③ r ≥ 1 なら移る。r < 1 なら確率 r で移る",
    ("", "移らなければ、同じ点をもう一度記録"),
    "④ ①〜③ を何千回もくり返す",
    "歩いた跡のヒストグラム = 事後分布",
    ("", "下がる方向にも確率的に進むので、山の広がり（誤差）まで分かる"),
], "③ がポイント: 上り坂 (もっともらしさが上がる) には必ず進み、下り坂にも r の確率で進む。いつも上るだけだと山頂 (最尤推定) で止まってしまい、誤差が分からない。\n"
   "このルールで歩くと、長い目で見て「各点にいる時間の割合」が事後分布 P(θ | D) に比例するようになる (詳細つり合い)。\n"
   "図は説明用のおもちゃの例 (実データではない): 1 つの倍率 f の事後分布を平均 1.48・幅 0.075 の山と仮定して 3000 歩。わざと山から離れた 1.2 から出発させたので、最初の 200 歩は山にたどり着くまでの途中 → 捨てる (バーンイン)。残りのヒストグラムが本当の分布 (点線) と重なる。\n"
   "図は slides/figs/plot_mcmc_toy.py で作った。",
   ref="N. Metropolis et al. (1953), J. Chem. Phys. 21, 1087", size=20, body_w=5.9)
s.shapes.add_picture("slides/figs/fig_mcmc_toy.png", Inches(6.9), Inches(2.0), width=Inches(6.2))
new.append(s)

new.append(slide("MCMC の中身③：今回の設定（emcee）", [
    "32 人の歩行者が同時に歩く。次の一歩は、他の歩行者の位置を使って決める",
    ("", "候補：Y = X_j + Z (X_k − X_j)（X_k：自分、X_j：ランダムに選んだ他の歩行者）"),
    ("", "Z は 1/2〜2 の乱数（確率 ∝ 1/√Z）、移る確率 = min(1, Z⁸ P(Y | D) / P(X_k | D))"),
    "出発点：最尤推定の点のすぐ近く。最初の 400 歩は捨て、5 歩ごとに記録",
    "結果の例（20.8 GeV）：ハローの倍率 f_halo = 1.48（68% 区間 1.40–1.55）",
    "注意：推奨の長さ（自己相関時間 τ の 50 倍）に届かないビンが多い",
    ("", "20.8 GeV で τ ≈ 127 歩、記録した長さ 5600 歩 < 50τ。13 ビン中 2 ビンだけ到達"),
    ("", "→ 誤差棒は目安。有意度 σ は Δln L から出すので影響しない"),
], "ストレッチ移動 (Goodman & Weare 2010): 自分 X_k と、他の歩行者 X_j を結ぶ直線の上で、Z 倍だけ伸ばしたり縮めたりした点を候補にする。他の歩行者の散らばり方を使うので、細長い山 (パラメータ同士が強く相関している、例えばガスと ICS の縮退) でも効率よく歩ける。\n"
   "Z⁸ の 8 は「パラメータの数 9 − 1」。直線上で伸び縮みさせた分の補正。Z の範囲 1/2〜2 は emcee の既定値 (a = 2)。\n"
   "68% 区間 = 記録したサンプルの下から 16% と 84% の値。\n"
   "自己相関時間 τ: 何歩離れたら前の位置を「忘れる」か。50τ 以上は emcee の推奨。足りない場合は最大 6000 歩まで自動で延ばしたが、それでも届かないビンが多い。\n"
   "数字は results/mcmc_allbins_gasICS_v20r_rerun/mcmc_bin06.json から。",
   ref="J. Goodman & J. Weare (2010), Comm. App. Math. Comp. Sci. 5, 65", size=20))

# 末尾に足した 3 枚を、MCMC のスライドの直後へ移す
lst = prs.slides._sldIdLst
ids = list(lst)
for k, sid in enumerate(ids[-3:]):
    lst.remove(sid)
    lst.insert(anchor + 1 + k, sid)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "枚 (挿入位置:", anchor + 2, "〜", anchor + 4, "枚目)")
