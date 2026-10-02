"""_claude 版: 捨てる歩数を 400 → 2000 に (2026-10-02 本人了承)。誤差は cloud_reports/2026-10-02_totani_mode/mw_burn2000/ の値。
実行: python slides/figs/edit_burn2000.py <入力 pptx> <出力 pptx>
"""
import sys
from pptx import Presentation

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
REPL = [
    ("最初の 400 歩（ふもとから登っている途中）は捨て", "最初の 2000 歩（ふもとから登っている途中）は捨て"),
    ("1.43〜1.61", "1.42〜1.58"),
    ("(32 人・6000 歩・400 歩)", "(32 人・6000 歩・最初の 2000 歩を捨てる)"),
]
NOTE_OLD = "400 歩は本研究のコードで決めた値"
NOTE_ADD = ("捨てる歩数は本研究で決めた値 (Totani の論文には書かれていない)。最初は 400 歩だったが、"
            "歩く人数 (20/32/64)・乱数の種 (3 通り)・歩数 (6000/20000) を変えて確かめたところ、"
            "400 歩では誤差の上端が乱数の種によって 1.59〜1.89 とぶれたので、2000 歩に変えた (2000 歩ではどれも 1.56〜1.57。この確認は型紙の作り方が少し違う確認用の計算で、本番の 2000 歩の誤差は 1.42〜1.58)。"
            "最良値 1.51 と有意度 19.5σ は捨てる歩数によらず同じ。400 歩のときの誤差は 1.43〜1.61。")
n = 0
for s in prs.slides:
    frames = [sh.text_frame for sh in s.shapes if sh.has_text_frame]
    if s.has_notes_slide:
        frames.append(s.notes_slide.notes_text_frame)
    for tf in frames:
        for p in tf.paragraphs:
            for r in p.runs:
                for a, b in REPL:
                    if a in r.text:
                        r.text = r.text.replace(a, b)
                        n += 1
    if s.has_notes_slide:
        nf = s.notes_slide.notes_text_frame
        for p in nf.paragraphs:
            t = "".join(r.text for r in p.runs)
            if t.startswith(NOTE_OLD):
                p.runs[0].text = NOTE_ADD
                for r in p.runs[1:]:
                    r.text = ""
                n += 1
prs.save(OUT)
print("saved", OUT, "置換", n)
