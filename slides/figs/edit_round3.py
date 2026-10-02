"""_claude 版の修正 (2026-10-02 本人の指摘 3 回目)。
実行: python slides/figs/edit_round3.py <入力 pptx> <出力 pptx>
"""
import copy
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
L = {l.name: l for l in prs.slide_layouts}
REF = next(sh for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Reference: D. Foreman"))


def find(prefix):
    return next(s for s in prs.slides if s.shapes.title is not None and s.shapes.title.text.startswith(prefix))


def index(prefix):
    return next(i for i, s in enumerate(prs.slides) if s.shapes.title is not None and s.shapes.title.text.startswith(prefix))


def body(s):
    return next(sh for sh in s.shapes if sh.has_text_frame and sh.shape_id != s.shapes.title.shape_id
                and not sh.text_frame.text.startswith("Reference"))


def para_set(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def replace_para(s, old, new):
    for p in body(s).text_frame.paragraphs:
        if "".join(r.text for r in p.runs) == old:
            para_set(p, new)
            return
    raise KeyError(old)


def insert_after(s, after_text, text, level, size):
    """after_text の段落の直後に、同じ書式の段落を足す"""
    tf = body(s).text_frame
    src = next(p for p in tf.paragraphs if "".join(r.text for r in p.runs) == after_text)
    new = copy.deepcopy(src._p)
    src._p.addnext(new)
    p = next(q for q in tf.paragraphs if q._p is new)
    para_set(p, text)
    p.level = level
    for r in p.runs:
        r.font.size = Pt(size)


def delete_para(s, text):
    tf = body(s).text_frame
    p = next(p for p in tf.paragraphs if "".join(r.text for r in p.runs) == text)
    p._p.getparent().remove(p._p)


def add_notes(s, text):
    nf = s.notes_slide.notes_text_frame
    nf.text = nf.text + "\n" + text


def new_slide(title, items, note, ref=None, width=11.8, size=22):
    s = prs.slides.add_slide(L["Title and Content"])
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(30)
    b = s.placeholders[1]
    b.left, b.top, b.width, b.height = Inches(0.6), Inches(1.5), Inches(width), Inches(4.7)
    tf = b.text_frame
    for k, it in enumerate(items):
        lvl, t = (1, it[1]) if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.text, p.level = t, lvl
        for r in p.runs:
            r.font.size = Pt(size - 4 if lvl else size)
    if ref:
        el = copy.deepcopy(REF._element)
        s.shapes._spTree.append(el)
        p = s.shapes[-1].text_frame.paragraphs[0]
        p.runs[1].text = " " + ref
        for r in p.runs[2:]:
            r._r.getparent().remove(r._r)
    s.notes_slide.notes_text_frame.text = note
    return s


def move_last_to(pos):
    lst = prs.slides._sldIdLst
    sid = list(lst)[-1]
    lst.remove(sid)
    lst.insert(pos, sid)


# 1. 「係数 9 個を決めるだけ」(手法② の前)
new_slide("ここからやること：係数 9 個を決めるだけ", [
    "予測 ＝ f₁ × 型紙₁ ＋ f₂ × 型紙₂ ＋ … ＋ f₉ × 型紙₉",
    ("", "型紙は固定。動かすのは係数（倍率）9 個だけ"),
    "尤度：その係数 9 個が「どれだけ良いか」の点数",
    "MCMC：良い係数 9 個を探す方法",
    "最良値：見つかった一番良い係数 9 個",
    "誤差：その係数が「どこまでずれても十分良いか」",
], "この後の尤度・有意度・MCMC のスライドは、全部「係数 9 個の決め方」の説明。難しい言葉が出てきたら、このスライドに戻って「今どの行の話か」を確認する。\n"
   "型紙 = 前の 7 枚で見た地図。係数 = その地図を何倍して足すか。", size=26)
move_last_to(index("手法②"))

# 2. Wilks
s = find("手法③")
replace_para(s, "根拠：Wilks の定理", "根拠：Wilks の定理（ガンマ線天文学の標準の換算。Totani の 13–19σ も同じ）")
add_notes(s, "Totani との関係: Totani (2025) Fig. 9 はハローあり・なしの ln L の差 (Δln L) を示しており、本文の 13–19σ はそれを σ = √(2Δln L) で σ に直した値。Totani は「Wilks の定理」という言葉は書いていないが、この換算はガンマ線天文学で標準 (2Δln L は TS = test statistic と呼ばれ、EGRET の解析 Mattox et al. 1996, ApJ 461, 396 から使われている)。本研究は Totani と同じ換算を使っている。")

# 3. MCMC の歩き方の種類 (④ の前)
new_slide("MCMC の歩き方はいろいろある", [
    "メトロポリス法（1953）：一番基本。少しずらして、比 r で移るか決める（②）",
    "ギブス法：倍率を 1 個ずつ順番に動かす",
    "ハミルトニアン MC：坂の傾きを使って、ボールを転がすように大きく動く",
    "アンサンブル法（2010）：何人も同時に歩き、仲間の位置で一歩を決める（④）",
    ("", "Python の道具 emcee がこれ。本研究はこれを使った"),
    "似ているが別物：焼きなまし法",
    ("", "下りに進む確率をだんだん小さくして、頂上だけを探す。誤差は出ない"),
    "Totani は「MCMC」とだけ書いていて、どの歩き方かは書いていない",
    ("", "本研究は、天文学で広く使われ、細長い山（ガスと ICS の縮退）に強い emcee を選んだ"),
], "どの歩き方でも、正しく長く歩けば足あとの分布は同じ山の形 (事後分布) になる。違うのは、歩く速さ (効率) と、どんな山が得意か。\n"
   "ギブス法: 9 個の倍率のうち 1 個だけを動かし、残り 8 個は止めておく。これを順番にくり返す。\n"
   "ハミルトニアン MC (HMC): 山の傾き (尤度の微分) を使って、遠くまで一気に動く。パラメータが多いときに強い。\n"
   "焼きなまし法: 金属をゆっくり冷やす「焼きなまし」が名前の由来。最初は下りにもよく進み、だんだん進みにくくして、最後は頂上で止まる。最良値を探す道具で、誤差は出ない。\n"
   "歩けているかの確認 (歩き方の種類ではない): 受理率 (足を出したうち、実際に移った割合。emcee では 0.2〜0.5 くらいが目安。今回の 20.8 GeV は 0.42)、自己相関時間 τ (何歩で前の場所を忘れるか)。\n"
   "emcee: Foreman-Mackey et al. (2013) の Python ライブラリ。中身は Goodman & Weare (2010) のアンサンブル法 (ストレッチ移動)。歩く人数・歩数・捨てる歩数は Totani の論文に書かれていないので、本研究で決めた (32 人・6000 歩・400 歩)。",
   ref="D. Foreman-Mackey et al. (2013), PASP 125, 306", size=20)
move_last_to(index("④ 32 人で歩く"))

# 4. ④ と ⑤
s = find("④ 32 人で歩く")
insert_after(s, "出発点：全員、Totani と同じ初期値の近く。6000 歩",
             "emcee はこの歩き方の Python の道具の名前。32 人・6000 歩は本研究の設定", 0, 18)
s = find("⑤ 歩き終わったら")
replace_para(s, "最初の 400 歩は捨て、下から 16%〜84% の範囲",
             "最初の 400 歩（ふもとから登っている途中）は捨て、残りの下から 16%〜84%")
add_notes(s, "なぜ最初を捨てるか: 出発点は Totani と同じ初期値で、山のふもとにある。最初のうちは頂上へ登っている途中なので、その足あとは「山の形」を表していない。だから誤差の計算からは外す (バーンイン)。一番高い足あと (最良値) は、捨てた分も含めて全部の中から探している。\n"
             "400 歩は本研究のコードで決めた値 (Totani の論文には書かれていない)。自己相関時間 τ ≈ 650 歩より短いので、誤差を出すには捨てる歩数を増やすほうが安全。最良値と有意度には影響しない。")

# 5. 電球のたとえ
s = find("問題：ハローと ICS")
insert_after(s, "ICS の倍率 1.39 → 0.48（20.8 GeV）",
             "たとえ：部屋の電球（ICS）と探している光（ハロー）が、写真ではほぼ同じ形 → 余った明るさがどちらのものか決めにくい", 0, 18)
s = find("① を調べる：手がかり")
insert_after(s, "→ 19 → 5σ は「手がかりを捨てた」から。汚染の証拠ではない",
             "たとえ：写真の真ん中を隠すと、2 つの光を見分けるヒントが減って自信が下がる", 0, 20)
s = find("② を調べる：σ を信じてよいか")
insert_after(s, "何も無い空 23 か所で、天の川と同じ解析をした",
             "たとえ：真っ暗な所を写したのに「明かりがある」と出てしまう（誤検出）が、どれくらい起きるか", 1, 17)

# 6. 疑問を 2 つに、③ のスライドを消す
s = find("縮退が引き起こした 3 つの疑問")
s.shapes.title.text = "縮退が引き起こした 2 つの疑問"
for r in s.shapes.title.text_frame.paragraphs[0].runs:
    r.font.size = Pt(30)
delete_para(s, "③ 7/23 に出した 19.1σ が再現できない")
delete_para(s, "結果そのものを信じてよいのか？")
nf = s.notes_slide.notes_text_frame
nf.text = "\n".join(l for l in nf.text.split("\n") if not l.startswith("(c)") and not l.startswith("③"))
i3 = index("③ を調べる：19.1σ")
lst = prs.slides._sldIdLst
sid = list(lst)[i3]
prs.part.drop_rel(sid.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
lst.remove(sid)

prs.save(OUT)
print("saved", OUT, len(prs.slides))
