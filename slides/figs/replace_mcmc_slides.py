"""_claude 版の「MCMC の中身①〜③」(3 枚) を、霧の山のたとえと対応させた 6 枚に置き換える。
文章は左半分だけに置き、右半分は本人がコードのスクショを貼れるように空けておく。
実行: python slides/figs/replace_mcmc_slides.py <入力 pptx> <出力 pptx>
"""
import copy
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
L = {l.name: l for l in prs.slide_layouts}
slides = list(prs.slides)
REF = next(sh for s in slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Reference: D. Foreman"))
old = [i for i, s in enumerate(slides) if s.shapes.title is not None and s.shapes.title.text.startswith("MCMC の中身")]
pos0 = old[0]

CODE = "コードは code/mcmc_fit_all_bins.py (Totani 方式の修正 cloud_reports/2026-10-02_totani_bestfit.patch を当てた後の行番号)。"


def add_ref(s, text):
    el = copy.deepcopy(REF._element)
    s.shapes._spTree.append(el)
    p = s.shapes[-1].text_frame.paragraphs[0]
    p.runs[1].text = " " + text
    for r in p.runs[2:]:
        r._r.getparent().remove(r._r)


def slide(title, items, note, ref=None, size=18):
    s = prs.slides.add_slide(L["Title and Content"])
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(30)
    b = s.placeholders[1]
    b.left, b.top, b.width, b.height = Inches(0.6), Inches(1.5), Inches(6.2), Inches(4.7)  # 右半分は空ける
    tf = b.text_frame
    for k, it in enumerate(items):
        lvl, t = (1, it[1]) if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.text, p.level = t, lvl
        for r in p.runs:
            r.font.size = Pt(size - 3 if lvl else size)
    if ref:
        add_ref(s, ref)
    s.notes_slide.notes_text_frame.text = note
    return s


new = []
new.append(slide("MCMC をたとえで見る：霧の山を歩く", [
    "山の上の「場所」 → 倍率の組 θ",
    ("", "θ = (f_iso, f_gas, f_ics, …, f_halo)。9 個あるので 9 次元の山"),
    "その場所の「高さ」 → 点数（事後分布）",
    "「高度計」 → 尤度の計算（手法② の ln L）",
    "「立ち入り禁止の場所」 → 事前分布",
    ("", "例：ガスや ICS の倍率は負にならない"),
    "「歩く人」 → 歩行者（walker）。今回は 32 人",
    "「足あとの記録」 → チェーン（サンプル）",
    "「一番高い足あと」 → 最良値",
    "「足あとの広がり」 → 誤差",
], "このスライドが「対応表」。この後のスライドは全部この言葉で説明する。\n"
   "9 次元の山は紙に描けない。だから地図全体を調べるのではなく、歩いて足あとを残し、足あとから山の形を知る。\n"
   "スクショ候補: 点数 (高さ) を計算する関数 log_probability (653〜659 行)。" + CODE))

new.append(slide("① 高さ＝点数の決め方（ベイズの定理）", [
    "P(θ | D) = L(D | θ) × π(θ) ÷ P(D)",
    "P(θ | D)：場所 θ の「高さ」（＝点数）",
    ("", "データ D を見た後で、θ がどれくらいもっともらしいか"),
    "L(D | θ)：「高度計」の読み（尤度）",
    ("", "倍率が θ のとき、観測した光子の地図 D がどれくらい出やすいか"),
    "π(θ)：入ってよい場所か（事前分布）",
    ("", "入れる場所は全部同じ値、立ち入り禁止なら 0"),
    "P(D)：海抜の基準。どの場所でも同じ数",
    ("", "「どちらが高いか」を比べるだけなら消える → 計算しなくてよい"),
    "計算では対数で：ln(高さ) = ln L + ln π",
], "ベイズの定理の読み方: 「データを見る前の考え (π: どこに入ってよいか)」に「データの当てはまり (L: 高度計)」を掛けると「データを見た後の考え (高さ)」になる。\n"
   "今回の π は一様: 入ってよい範囲 (ガス・ICS などは 0 以上、全部 |f| ≤ 10⁶) なら同じ値、外なら 0。だから高さの形は高度計の読み (尤度) と同じになる。\n"
   "P(D) は「海抜何メートルから測るか」の基準のようなもの。2 地点の高さを比べるときは、どちらにも同じ基準が入っているので割り算で消える。MCMC は比べることしかしないので、P(D) を計算しなくてよい。\n"
   "対数にするのは、数が大きすぎたり小さすぎたりしてコンピュータで扱えなくなるのを防ぐため。\n"
   "スクショ候補: log_prior (630 行、立ち入り禁止のチェック)、log_likelihood (647〜650 行、高度計)、log_probability (653 行、両方を足す)。" + CODE,
   ref="D. Foreman-Mackey et al. (2013), PASP 125, 306"))

new.append(slide("② 一歩の進み方（メトロポリス法）", [
    "1. 今いる場所 θ から、少しずれた場所 θ′ に足を出す",
    ("", "どれだけずらすかは乱数で決める"),
    "2. 高度計で両方を読み、比 r = 高さ(θ′) ÷ 高さ(θ)",
    "3. r ≥ 1（上り）なら、必ず θ′ に移る",
    "4. r < 1（下り）なら、確率 r で移る",
    ("", "0〜1 の乱数を引き、r より小さければ移る"),
    "5. 移っても移らなくても、今いる場所を足あとに記録",
    "6. 1〜5 を何千回もくり返す",
    "なぜ下りにも進む？",
    ("", "上りだけだと頂上で止まり、山の広がり（誤差）が分からないから"),
], "例: 今いる場所の高さが 100、足を出した先の高さが 80 なら r = 0.8。0〜1 の乱数を引いて 0.8 より小さければ (80% の確率で) 移る。高さが 10 なら r = 0.1 で、10% しか移らない。だから、低い場所にはめったに行かず、高い場所に長くいる。\n"
   "長く歩くと「各場所にいた時間の割合」が「その場所の高さ」に比例するようになる (詳細つり合い)。だから足あとのヒストグラムが山の形 (事後分布) になる。\n"
   "移らなかったときも今の場所をもう一度記録するのが大事。そうしないと、高い場所にとどまった回数が数えられない。\n"
   "スクショ候補: 歩かせる部分 sampler.run_mcmc (877〜878 行)。メトロポリスの判定そのものは emcee の中にある。" + CODE,
   ref="N. Metropolis et al. (1953), J. Chem. Phys. 21, 1087"))

s = prs.slides.add_slide(L["Title Only"])
s.shapes.title.text = "③ 足あとはこうなる（説明用の例）"
for r in s.shapes.title.text_frame.paragraphs[0].runs:
    r.font.size = Pt(30)
s.shapes.add_picture("slides/figs/fig_mcmc_toy.png", Inches(0.8), Inches(1.4), width=Inches(11.7))
tb = s.shapes.add_textbox(Inches(0.6), Inches(6.2), Inches(12.2), Inches(0.9))
tb.text_frame.word_wrap = True
tb.text_frame.text = "左：足あと（縦＝倍率、横＝歩数）。最初の 200 歩は山にたどり着く途中なので捨てる。右：足あとのヒストグラム ＝ 山の形"
for r in tb.text_frame.paragraphs[0].runs:
    r.font.size = Pt(18)
s.notes_slide.notes_text_frame.text = (
    "実データではない、説明用のおもちゃの例。倍率が 1 個だけの山 (1 次元) を、平均 1.48・幅 0.075 の形だと仮定して 3000 歩歩かせた。\n"
    "わざと山から離れた 1.2 から出発させたので、最初はふもとから登っていく。この登っている途中 (最初の 200 歩) は山の形と関係ないので捨てる (バーンイン)。\n"
    "そのあとは山の上を行ったり来たりしている。足あとのヒストグラム (右の灰色) が、本当の山の形 (点線) と重なっている = 足あとから山の形が分かる。\n"
    "図は slides/figs/plot_mcmc_toy.py で作った (コメント付き)。")
new.append(s)

new.append(slide("④ 32 人で歩く（emcee）", [
    "1 人だと、細長い山で迷いやすい",
    ("", "例：ガスと ICS の縮退。片方を上げて片方を下げても高さがほぼ同じ"),
    "そこで 32 人が同時に歩き、仲間の位置を見て一歩を決める",
    "一歩の出し方（ストレッチ移動）",
    ("", "自分 X_k から、ランダムに選んだ仲間 X_j への線を引く"),
    ("", "その線の上で、Z 倍（1/2〜2 の乱数）伸ばした点に足を出す"),
    ("", "Y = X_j + Z (X_k − X_j)"),
    "移るかどうかは ② と同じ考え（Z⁸ を掛けて補正）",
    "出発点：全員、Totani と同じ初期値の近く。6000 歩",
], "細長い山: 仲間たちの散らばり方自体が「山が細長い方向」を表している。仲間のいる方向に沿って一歩を出すので、細長い山でも効率よく歩ける。\n"
   "Z⁸ の 8 は「倍率の数 9 − 1」。線の上で伸び縮みさせた分、足を出せる場所の広さが変わるので、それを補正する。Z の範囲 1/2〜2 は emcee の決まった値 (a = 2)。\n"
   "Totani と同じ初期値: 点源・GALPROP (ガス・ICS) = 1 倍、等方 = E²dN/dE 1e-4 相当、Loop I・バブル・ハロー = 0 (Totani §2.3)。\n"
   "スクショ候補: 出発点を決める行 (868 行)、32 人で歩かせる emcee.EnsembleSampler (877 行)。" + CODE,
   ref="J. Goodman & J. Weare (2010), Comm. App. Math. Comp. Sci. 5, 65"))

new.append(slide("⑤ 歩き終わったら：最良値と誤差", [
    "一番高い足あと → 最良の倍率（Totani と同じ）",
    ("", "20.8 GeV：ハローの倍率 1.51"),
    "足あとの広がり → 誤差",
    ("", "最初の 400 歩は捨て、下から 16%〜84% の範囲"),
    ("", "20.8 GeV：1.43〜1.61"),
    "有意度：ハローあり・なしの山を別々に歩く",
    ("", "それぞれの一番高い足あとの高さの差 Δln L → σ = √(2Δln L)"),
    ("", "20.8 GeV：19.5σ"),
    "注意：山の端まで歩き回るには、もっと長く歩くのが理想",
    ("", "誤差は目安。一番高い足あとは頂上に届いている"),
], "ハローなしの山: ハローの倍率を 0 に固定した 8 次元の山。ハローありの山 (9 次元) より低いか同じ。2 つの頂上の高さの差が Δln L。\n"
   "頂上に届いているかの確認: 坂を登る方法 (勾配法) で頂上を求めて、一番高い足あととの差を調べた。13 ビンすべてで ln L の差 0.6 以下、有意度の差 0.28σ 以下。確認のためだけで、結果には使っていない。\n"
   "長さの目安: 自己相関時間 τ (何歩離れたら前の場所を忘れるか) の 50 倍以上が推奨。20.8 GeV で τ ≈ 650 歩なので 3 万歩以上が理想だが、6000 歩。\n"
   "数字はクラウドの計算機で、Totani と同じやり方で出した値 (cloud_reports/2026-10-02_totani_mode/)。\n"
   "スクショ候補: 一番高い足あとを探す行 (879〜880 行)、誤差の 16〜84% (926 行)、有意度 (910 行)。" + CODE,
   ref="T. Totani (2025), arXiv:2507.07209 §2.2"))

# 古い 3 枚を消して、新しい 6 枚をその位置へ
lst = prs.slides._sldIdLst
ids = list(lst)
for i in sorted(old, reverse=True):
    sid = ids[i]
    prs.part.drop_rel(sid.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
    lst.remove(sid)
ids = list(lst)
for k, sid in enumerate(ids[-len(new):]):
    lst.remove(sid)
    lst.insert(pos0 + k, sid)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "枚。新しい 6 枚は", pos0 + 1, "〜", pos0 + 6, "枚目")
