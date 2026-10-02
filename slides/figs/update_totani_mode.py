"""_claude 版の数字・図・解説を、Totani (2025) と同じ最良値の求め方 (MCMC で踏んだ点のうち尤度が最大の点) で
計算し直した結果 (2026-10-02、クラウド) に差し替える。
実行: python slides/figs/update_totani_mode.py <入力 pptx> <出力 pptx>
"""
import sys

from pptx import Presentation
from pptx.util import Emu

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
FIG = "slides/figs/"

# 本文の置き換え (段落の中の文字列を置き換える)
BODY = [
    # 手法④ (MCMC)
    ("手法④：MCMC（誤差の見積もり）", "手法④：MCMC（最良値と誤差）"),
    ("最尤の 1 点だけでなく、倍率 f が「どれくらいの範囲にありうるか」を知りたい",
     "MCMC で、最良の倍率 f と、その誤差の両方を求める（Totani と同じ）"),
    ("歩いた跡の分布 = パラメータの事後分布 → 誤差棒",
     "歩いた中で尤度が最大の点 = 最良値、歩いた跡の分布 → 誤差棒"),
    ("emcee を使用：32 本の歩行者、最初の 400 歩は捨てる", "emcee を使用：32 本の歩行者 × 6000 歩"),
    ("自己相関時間の 50 倍以上の長さがあるか確認", "出発点は Totani と同じ初期値（点源・GALPROP = 1、他 = 0 など）"),
    ("有意度 σ は Δln L から出すので、MCMC の有無で変わらない",
     "有意度 σ も、ハローあり・なしそれぞれの MCMC で見つけた最大 ln L の差から出す"),
    # MCMC の中身③
    ("出発点：最尤推定の点のすぐ近く。最初の 400 歩は捨て、5 歩ごとに記録",
     "出発点：Totani と同じ初期値。6000 歩歩き、尤度が最大の点を最良値にする"),
    ("結果の例（20.8 GeV）：ハローの倍率 f_halo = 1.48（68% 区間 1.40–1.55）",
     "結果の例（20.8 GeV）：ハローの倍率 f_halo = 1.51（68% 区間 1.43–1.61）"),
    ("20.8 GeV で τ ≈ 127 歩、記録した長さ 5600 歩 < 50τ。13 ビン中 2 ビンだけ到達",
     "20.8 GeV で τ ≈ 650 歩、6000 歩 < 50τ"),
    ("→ 誤差棒は目安。有意度 σ は Δln L から出すので影響しない",
     "→ 誤差棒は目安。ただし最良値は尤度の頂上に届いている（頂上との差は ln L で 0.6 以下）"),
    # 結果: 再現
    ("20.8 GeV で 19.0σ（Totani は 13–19σ）。山の位置も一致", "20.8 GeV で 19.5σ（Totani は 13–19σ）。山の位置も一致"),
    # 縮退
    ("ハローの光の 8〜10 割は ICS から移ったもの", "ICS の倍率 1.39 → 0.48（20.8 GeV）"),
    ("① バブル領域を外すと 19.0σ → 5.0σ に落ちる", "① バブル領域を外すと 19.5σ → 5.0σ に落ちる"),
    ("手がかりの減り方だけで予測すると 6.4σ、実測 5.0σ", "手がかりの減り方だけで予測すると 6.6σ、実測 5.0σ"),
    ("同じ形の領域を横にずらして外すと 13〜16σ（予測とも一致）", "同じ形の領域を横にずらして外すと 15〜16σ"),
    ("19.0σ を較正すると 19.0 ÷ 1.26 ≈ 15σ（13.7–16.6σ）", "19.5σ を較正すると 19.5 ÷ 1.26 ≈ 15σ（14.0–17.0σ）"),
    ("→ 主結果を 19.0σ（10/1 の再計算）に統一", "→ 10/2 に Totani と同じやり方で計算し直し、主結果は 19.5σ"),
    # 矮小銀河
    ("暗黒物質があれば：T の平均 = √( Σ pred_i² ) = 1.73", "暗黒物質があれば：T の平均 = √( Σ pred_i² ) = 1.77"),
    ("実測 T = 0.39 → ずれ 1.3σ", "実測 T = 0.39 → ずれ 1.4σ"),
    ("誤差なし：E[T] = Σ w_i pred_i / √Σ w_i² = 1.73", "誤差なし：E[T] = Σ w_i pred_i / √Σ w_i² = 1.77"),
    ("期待 1.7σ → J の誤差で 1.3〜5.0σ に広がる。実測 0.4σ ± 1.26。ずれ 1.3σ",
     "期待 1.8σ → J の誤差で 1.4〜5.1σ に広がる。実測 0.4σ ± 1.26。ずれ 1.3σ"),
    ("J の誤差なし：期待 1.73、ずれ 1.3σ（s = 1.26 なら 1.1σ）", "J の誤差なし：期待 1.77、ずれ 1.4σ（s = 1.26 なら 1.1σ）"),
    ("J の誤差あり：期待 2.4（68% で 1.3–5.0）、ずれ 1.5σ（1.3σ）", "J の誤差あり：期待 2.5（68% で 1.4–5.1）、ずれ 1.5σ（1.3σ）"),
    # 15σ と 1.3σ、結論
    ("19.0σ を、何も無い空でのばらつき 1.26 で割ったもの", "19.5σ を、何も無い空でのばらつき 1.26 で割ったもの"),
    ("Totani の超過を再現（名目 19.0σ、較正後 約 15σ）", "Totani の超過を再現（名目 19.5σ、較正後 約 15σ）"),
]

# ノートの置き換え
NOTE = [
    ("最大化は L-BFGS-B (範囲の制約つきの勾配法) で、出発点を変えて何度か行う。",
     "最大化は Totani (2025) §2.2 と同じく MCMC (emcee) で行い、歩いた中で ln L が最大の点を最良値とする。"),
    ("数字は results/mcmc_allbins_gasICS_v20r_rerun/mcmc_bin06.json から。",
     "数字は 10/2 に Totani と同じやり方で計算し直した結果 (クラウド、cloud_reports/2026-10-02_totani_mode/)。最良値が頂上に届いているかは、勾配法で頂上を求めて確認した (13 ビンすべてで ln L の差 0.6 以下、有意度の差 0.28σ 以下。確認のためだけで、結果には使っていない)。"),
    ("青が本研究 (v20r、10/1 にコードの版と設定の記録つきで再計算)、赤が Totani Fig. 9。",
     "青が本研究 (10/2 に Totani と同じやり方 = MCMC で最良値、で計算し直した。クラウドの計算機)、赤が Totani Fig. 9。本人 PC での確定値は後日。"),
    ("ただしこの 19.0σ は Wilks の定理を信じた「名目の値」。", "ただしこの 19.5σ は Wilks の定理を信じた「名目の値」。"),
    ("ハローなしでは全エネルギーで約 1.6 倍で平ら。", "ハローなしでは全エネルギーで約 1.4〜1.6 倍でほぼ平ら (最良値)。"),
    ("予測 6.4σ と実測 5.0σ", "予測 6.6σ と実測 5.0σ"),
    ("ずらして外すと 13.4〜16.3σ。", "ずらして外すと 14.6〜16.4σ (Totani と同じやり方)。"),
    ("19.5σ と 18.4σ はクラウドの計算機で同じ条件で比べた値 (本人 PC の基準は 19.0σ なので、",
     "19.5σ と 18.4σ はクラウドの計算機で、Totani と同じやり方で比べた値 (どちらも同じ条件なので、"),
    ("有意度は 19.1σ → 19.0σ でほぼ同じ。", "有意度は 19.1σ → 19.0σ でほぼ同じ。10/2 に最良値の求め方を Totani と同じ (MCMC) に直して計算し直し、19.5σ (クラウド)。"),
    ("(成分比: ガス 0.91 / ICS 1.15)", "(成分比 = 最良値の Totani 比: ガス 0.95 / ICS 1.05 / ハロー 1.21、Totani と同じやり方)"),
    ("ずれ = (1.73 − 0.39) ÷ 1 = 1.3σ。", "ずれ = (1.77 − 0.39) ÷ 1 = 1.4σ。天の川のハローの明るさは Totani と同じやり方の最良値。"),
    ("点線: J の誤差なしの予測 1.73。", "点線: J の誤差なしの予測 1.77。"),
    ("ρ☉ = 0.3: 期待 3.38 (誤差なし) / 4.74 (誤差あり、68% 2.58–9.72)、ずれ 2.4σ / 2.0σ (s = 1.26)。",
     "ρ☉ = 0.3: 期待 3.46 (誤差なし) / 4.86 (誤差あり、68% 2.64–9.95)、ずれ 2.4σ / 2.1σ (s = 1.26)。"),
    ("ρ☉ = 0.6: 期待 0.85 / 1.19 (0.64–2.43)、ずれ 0.4σ / 0.7σ。", "ρ☉ = 0.6: 期待 0.87 / 1.21 (0.66–2.49)、ずれ 0.4σ / 0.7σ。"),
]

# 図の差し替え (古い図の画素サイズ → 新しい図)
PICS = {
    (2175, 1140): FIG + "fig_bin6_before_after_totani.png",
}
SIG_NEW = FIG + "fig_significance_totani.png"
ICS_NEW = FIG + "fig_ics_norm_totani.png"
J_NEW = FIG + "fig_dwarf_J_uncertainty_totani.png"


def sub_paragraphs(tf, old, new):
    n = 0
    for p in tf.paragraphs:
        t = "".join(r.text for r in p.runs)
        if old in t and p.runs:
            p.runs[0].text = t.replace(old, new)
            for r in p.runs[1:]:
                r.text = ""
            n += 1
    return n


def swap_picture(slide, sh, path):
    l, t, w = sh.left, sh.top, sh.width
    crop = (sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom)
    h_old = sh.height
    new = slide.shapes.add_picture(path, l, t, width=w)
    new.crop_left, new.crop_right, new.crop_top, new.crop_bottom = crop
    if any(crop):
        new.height = h_old
    sh._element.getparent().replace(sh._element, new._element)


used = {o: 0 for o, _ in BODY}
used_n = {o: 0 for o, _ in NOTE}
for s in prs.slides:
    title = s.shapes.title.text if s.shapes.title is not None else ""
    for sh in list(s.shapes):
        if sh.has_text_frame:
            for o, n in BODY:
                used[o] += sub_paragraphs(sh.text_frame, o, n)
        if sh.shape_type == 13:
            size = sh.image.size
            if size in PICS:
                swap_picture(s, sh, PICS[size])
            elif title.startswith("結果：Totani の超過を再現"):
                swap_picture(s, sh, SIG_NEW)
            elif title.startswith("問題：ハローと ICS"):
                swap_picture(s, sh, ICS_NEW)
            elif title.startswith("結果：予測の幅と実測"):
                swap_picture(s, sh, J_NEW)
    if s.has_notes_slide:
        nf = s.notes_slide.notes_text_frame
        txt = nf.text
        for o, n in NOTE:
            if o in txt:
                txt = txt.replace(o, n)
                used_n[o] += 1
        if txt != nf.text:
            nf.text = txt

miss = [o for o, k in used.items() if k == 0] + [o for o, k in used_n.items() if k == 0]
print("置き換えできなかったもの:", miss if miss else "なし")
prs.save(OUT)
print("saved", OUT, len(prs.slides))
